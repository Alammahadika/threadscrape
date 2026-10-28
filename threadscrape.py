#!/usr/bin/env python3
"""
threadscrape.py - Threads search results (HTML or live) -> raw data -> .xlsx
Collection only: no text cleaning here. Do cleaning in a separate step.

Mode 1 (offline) : python threadscrape.py --html results.html --query "forest fire"
Mode 2 (live)    : python threadscrape.py --query "forest fire" --scrolls 15
Login once       : python threadscrape.py --login

Install:
    pip install beautifulsoup4 lxml pandas openpyxl playwright
    playwright install chromium

Selectors avoid hashed class names; they use semantic attributes
(href /@user/post/ID, <time datetime>, role, svg title).
"""
import argparse
import random
import re
import sys
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import pandas as pd
from bs4 import BeautifulSoup

STATE_FILE = "threads_state.json"

# icon title keywords (multi-language: id/en/ru) -> column name
METRIC_KEYS = {
    "likes": ("suka", "like", "нравится"),
    "replies": ("balas", "reply", "ответ", "komentar", "comment"),
    "reposts": ("posting ulang", "repost", "репост"),
    "shares": ("bagikan", "kirim", "share", "поделиться", "send"),
}
VERIFIED_KEYS = ("verified", "terverifikasi", "подтверждено")
POST_RE = re.compile(r"/@([^/]+)/post/([\w-]+)")
HASHTAG_RE = re.compile(r"#(\w+)")
MENTION_RE = re.compile(r"@\w[\w.]*")


# ----------------------------------------------------------------- parsing
def parse_count(raw: str):
    """'36 632' -> 36632 ; '1,2K' -> 1200 ; '3,4 rb' -> 3400 ; '2jt' -> 2000000"""
    if not raw:
        return 0
    s = raw.replace("\xa0", " ").strip().lower()
    m = re.match(r"^([\d.,\s]+)\s*(k|rb|ribu|m|jt|juta|tho|тыс)?\.?$", s)
    if not m:
        return 0
    num, suf = m.group(1).replace(" ", ""), m.group(2)
    mult = {"k": 1e3, "rb": 1e3, "ribu": 1e3, "tho": 1e3, "тыс": 1e3,
            "m": 1e6, "jt": 1e6, "juta": 1e6}.get(suf, 1)
    if mult == 1:  # plain number: drop thousands separators
        return int(re.sub(r"[.,]", "", num) or 0)
    return int(float(num.replace(",", ".")) * mult)


def extract_text(box):
    """Main post text = <span> direct child of span[dir=auto], not inside a link/button."""
    for sp in box.select('span[dir="auto"]'):
        if sp.find_parent("a") or sp.find_parent(attrs={"role": "button"}):
            continue
        child = sp.find("span", recursive=False)
        if not child:
            continue
        for br in child.find_all("br"):
            br.replace_with("\n")
        txt = child.get_text("")
        if txt.strip():
            return txt
    return ""


def parse_post(box):
    link = next((a for a in box.select('a[href*="/post/"]')
                 if POST_RE.search(a.get("href", "")) and not a["href"].endswith("/media")), None)
    if not link:
        return None
    user, pid = POST_RE.search(link["href"]).groups()

    t = box.find("time")
    ts = t.get("datetime") if t else None

    row = {
        "post_id": pid,
        "username": user,
        "url": f"https://www.threads.com/@{user}/post/{pid}",
        "timestamp_raw": ts,
        "text": extract_text(box),
        "topic_tag": "",
        "verified": False,
        "media": "",
        "likes": 0, "replies": 0, "reposts": 0, "shares": 0,
    }

    tag = box.select_one('a[href*="serp_type=tags"]')
    if tag:
        row["topic_tag"] = tag.get_text(strip=True)

    # verified badge + metrics via svg icon titles
    for svg in box.find_all("svg"):
        title = svg.find("title")
        title = (title.get_text() if title else svg.get("title", "")).lower()
        if any(k in title for k in VERIFIED_KEYS):
            row["verified"] = True
    for btn in box.select('div[role="button"]'):
        svg = btn.find("svg")
        if not svg:
            continue
        ti = svg.find("title")
        ti = (ti.get_text() if ti else svg.get("title", "")).lower()
        for col, keys in METRIC_KEYS.items():
            if any(k in ti for k in keys):
                nums = [s.get_text(strip=True) for s in btn.find_all("span")
                        if re.search(r"\d", s.get_text())]
                row[col] = parse_count(nums[-1]) if nums else 0
                break

    if box.find("video"):
        row["media"] = "video"
    elif box.select_one('a[href$="/media"]'):
        row["media"] = "image"
    return row


def parse_html(html: str):
    soup = BeautifulSoup(html, "lxml")
    out = []
    for box in soup.select('[data-pressable-container="true"]'):
        r = parse_post(box)
        if r:
            out.append(r)
    return out


# ----------------------------------------------------------------- dataframe
def build_df(rows, query):
    if not rows:
        return pd.DataFrame()
    df = pd.DataFrame(rows).drop_duplicates("post_id", keep="last").reset_index(drop=True)

    # minimal hygiene only (non-breaking spaces / edge whitespace); no content cleaning
    df["text"] = df["text"].fillna("").str.replace("\xa0", " ").str.strip()

    ts = pd.to_datetime(df["timestamp_raw"], utc=True, errors="coerce")
    df["timestamp"] = ts.dt.tz_localize(None)  # UTC, timezone-naive for Excel
    df["date"] = df["timestamp"].dt.date

    df["hashtags"] = df["text"].apply(lambda s: ", ".join(HASHTAG_RE.findall(s)))
    df["mentions"] = df["text"].apply(lambda s: ", ".join(MENTION_RE.findall(s)))
    df["engagement"] = df[["likes", "replies", "reposts", "shares"]].sum(axis=1)
    df["query"] = query

    # drop empty posts (no text and no media)
    df = df[(df["text"] != "") | (df["media"] != "")]
    cols = ["query", "post_id", "username", "verified", "timestamp", "date",
            "text", "hashtags", "mentions", "topic_tag", "media",
            "likes", "replies", "reposts", "shares", "engagement", "url"]
    return df[cols].sort_values("timestamp", ascending=False).reset_index(drop=True)


# ----------------------------------------------------------------- export
def export_xlsx(df, path, meta):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm:ss") as xw:
        df.to_excel(xw, sheet_name="posts", index=False)
        pd.DataFrame(list(meta.items()), columns=["field", "value"]).to_excel(
            xw, sheet_name="info", index=False)
        ws = xw.sheets["posts"]
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        widths = {"text": 80, "url": 45, "hashtags": 25, "mentions": 25, "timestamp": 20}
        for i, col in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 14)
        for r in ws.iter_rows(min_row=2):
            for c in r:
                if ws.cell(1, c.column).value == "text":
                    c.alignment = Alignment(wrap_text=True, vertical="top")


# ----------------------------------------------------------------- live (Playwright)
def login():
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch(headless=False)
        ctx = b.new_context(locale="en-US")
        pg = ctx.new_page()
        pg.goto("https://www.threads.com/login")
        input("Log in in the browser, then press ENTER here... ")
        ctx.storage_state(path=STATE_FILE)
        b.close()
    print(f"Session saved: {STATE_FILE}")


def scrape_live(query, scrolls, headless, patience=8, max_errors=5):
    from playwright.sync_api import sync_playwright
    from playwright.sync_api import TimeoutError as PWTimeout
    if not Path(STATE_FILE).exists():
        sys.exit("No saved session. Run first: python threadscrape.py --login")

    SEL = '[data-pressable-container="true"]'
    RETRY = re.compile(r"try again|retry|coba lagi|повторить", re.I)
    STATE_JS = "(s) => [document.querySelectorAll(s).length, document.documentElement.scrollHeight]"
    CHANGED_JS = ("([s, n, h]) => document.querySelectorAll(s).length !== n"
                  " || document.documentElement.scrollHeight > h")

    seen, rows, stale, errors = set(), [], 0, 0
    reason = f"reached {scrolls} scrolls"
    url = f"https://www.threads.com/search?q={quote(query)}&serp_type=default&filter=recent"

    with sync_playwright() as p:
        b = p.chromium.launch(headless=headless)
        try:
            ctx = b.new_context(storage_state=STATE_FILE, locale="en-US",
                                viewport={"width": 1280, "height": 900})
            pg = ctx.new_page()
            pg.goto(url, wait_until="domcontentloaded")
            pg.wait_for_selector(SEL, timeout=30000)

            for i in range(scrolls):
                try:
                    new = 0
                    for r in parse_html(pg.content()):
                        if r["post_id"] not in seen:
                            seen.add(r["post_id"])
                            rows.append(r)
                            new += 1
                    stale = stale + 1 if new == 0 else 0
                    errors = 0
                    print(f"scroll {i + 1}/{scrolls}: +{new} "
                          f"(total {len(rows)}, stale {stale}/{patience})")
                    if stale >= patience:
                        reason = f"no new posts for {patience} consecutive scrolls"
                        break

                    n, h = pg.evaluate(STATE_JS, SEL)
                    pg.evaluate("window.scrollTo(0, document.documentElement.scrollHeight)")
                    pg.mouse.wheel(0, random.randint(1500, 2500))
                    try:
                        pg.wait_for_function(CHANGED_JS, arg=[SEL, n, h], timeout=8000)
                    except PWTimeout:
                        btn = pg.get_by_role("button", name=RETRY)
                        if btn.count():
                            btn.first.click()
                            print("  clicked retry button")
                        pg.keyboard.press("End")
                        time.sleep(min(5 * (stale + 1), 30))
                    time.sleep(random.uniform(1.0, 2.0))

                except KeyboardInterrupt:
                    raise
                except Exception as e:  # DOM re-render, navigation, etc.
                    errors += 1
                    stale += 1
                    print(f"  step failed ({errors}/{max_errors}): {str(e).splitlines()[0]}")
                    if errors >= max_errors:
                        reason = f"{max_errors} consecutive errors"
                        break
                    time.sleep(3)

        except KeyboardInterrupt:
            reason = "interrupted by user (partial results kept)"
        finally:
            b.close()

    print(f"Stopped: {reason}")
    return rows

# ----------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="threadscrape: Threads search -> xlsx (raw collection)")
    ap.add_argument("--html", nargs="+", help="saved HTML file(s) of search results (offline mode)")
    ap.add_argument("--query", default="", help="2-3 word search query, e.g. 'forest fire'")
    ap.add_argument("--scrolls", type=int, default=10, help="number of scrolls (live mode)")
    ap.add_argument("--headed", action="store_true", help="show the browser (live mode)")
    ap.add_argument("--out", help="output .xlsx filename")
    ap.add_argument("--login", action="store_true", help="log in once and save the session")
    ap.add_argument("--patience", type=int, default=8,help="stop after N consecutive scrolls without new posts (live mode)")
    a = ap.parse_args()

    if a.login:
        return login()

    if a.html:
        rows = []
        for f in a.html:
            rows += parse_html(Path(f).read_text(encoding="utf-8", errors="ignore"))
        mode = "html"
    elif a.query:
        rows = scrape_live(a.query, a.scrolls, headless=not a.headed, patience=a.patience)
        mode = "live"
    else:
        ap.error("provide --html FILE or --query 'keywords'")

    df = build_df(rows, a.query)
    if df.empty:
        sys.exit("No posts parsed. Check that the HTML contains the full search results.")

    slug = re.sub(r"\W+", "_", a.query.strip().lower()) or "threads"
    out = a.out or f"threads_{slug}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    export_xlsx(df, out, {
        "query": a.query,
        "mode": mode,
        "total_posts": len(df),
        "timestamp_range_utc": f"{df['timestamp'].min()} to {df['timestamp'].max()}",
        "collected_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
    })
    print(f"OK: {len(df)} posts -> {out}")


if __name__ == "__main__":
    main()
    
